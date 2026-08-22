import openpyxl
import os
import sys
import time

# Ensure current folder is in path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from agentic_rag import AgenticRAGOrchestrator, LLM_MODEL, VECTOR_DB_DIR, MONITORING_DB
from agentic_app import sanitise_answer

xlsx_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Chatbot Eval Dataset.xlsx")
if not os.path.exists(xlsx_path):
    print(f"Error: {xlsx_path} not found!")
    sys.exit(1)

wb = openpyxl.load_workbook(xlsx_path)
sheet = wb["Golden Dataset"]

print("Initializing RAG Orchestrator ...")
orch = AgenticRAGOrchestrator(LLM_MODEL, VECTOR_DB_DIR, MONITORING_DB)

output_log = os.path.join(os.path.dirname(os.path.abspath(__file__)), "eval_suite_results.txt")
print(f"Writing results to: {output_log}")

questions_run = 0

with open(output_log, "w", encoding="utf-8") as f:
    f.write("=== RAG EVALUATION SUITE RUN ===\n\n")
    
    # Read rows 2 to 55 (54 questions)
    for row_idx in range(2, 56):
        q_id = sheet.cell(row_idx, 1).value
        q_type = sheet.cell(row_idx, 2).value
        q_text = sheet.cell(row_idx, 3).value
        gt_answer = sheet.cell(row_idx, 4).value
        
        if not q_text:
            continue
            
        print(f"Running Q{q_id} ({q_type}): {q_text[:60]}...")
        t0 = time.perf_counter()
        
        try:
            resp = orch.query_sync(q_text, session_id=f"eval-q{q_id}")
            clean_ans = sanitise_answer(resp.answer)
            elapsed = time.perf_counter() - t0
            
            f.write(f"QID: {q_id}\n")
            f.write(f"Type: {q_type}\n")
            f.write(f"Question: {q_text}\n")
            f.write(f"Ground Truth: {gt_answer}\n")
            f.write(f"Chatbot Answer: {clean_ans}\n")
            f.write(f"Confidence: {resp.confidence:.0%}\n")
            f.write(f"Time: {elapsed:.2f}s\n")
            f.write(f"Citations: {resp.citations.strip().replace('\n', ' | ')}\n")
            f.write("-" * 80 + "\n\n")
            f.flush()
            questions_run += 1
        except Exception as e:
            f.write(f"QID: {q_id} failed with error: {e}\n\n")
            f.flush()

print(f"\nDone. Successfully executed {questions_run} evaluation questions.")
print(f"Results saved in: {output_log}")
